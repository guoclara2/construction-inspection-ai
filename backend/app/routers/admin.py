# 管理端路由：知识库 / 项目 / 成员 / 工单监督 / 统计看板 / SLA 规则（全程按 project_scope 与权限矩阵隔离）
# P05（§2.3/2.4）：监督列表与看板数据库层分页/截断 + overdue SQL 化 + sla-rules 配置接口
import io
import json
from datetime import timedelta

from fastapi import APIRouter, Depends, File, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import func as sa_func, text, update
from sqlalchemy.orm import Session

from ..auth import hash_password, revoke_all_refresh
from ..database import get_db
from ..deps import ProjectScope, project_scope, require_perm
from ..models import (
    AiFeedback,
    InspectionItem,
    InspectionRecord,
    Notification,
    OrderLog,
    Project,
    ProjectMember,
    RectifyOrder,
    SlaRule,
    User,
)
from ..security import generate_strong_password
from ..services.audit import write_audit
from ..services.metrics import sla_on_time_rate
from ..services.notify import _send_one
from ..services.sla import resolve_urge_interval_hours
from ..utils import BizError, as_utc, fmt, forbidden, now_utc, ok
from .inspection import build_item_query, item_to_dict
from .orders import order_to_dict

router = APIRouter(prefix="/api/admin", tags=["admin"])

# 合法枚举集合
DEFECT_CATEGORIES = {"structural", "protection", "installation", "material", "safety"}
RISK_LEVELS = {"high", "medium", "low"}
PROJECT_TYPES = {"building", "road", "tunnel", "landscape"}
PHASES = {"foundation", "structure", "mep", "decoration"}
MEMBER_ROLES = {"inspector", "rectifier", "reviewer", "project_admin", "viewer"}
SPECIALTIES = {"civil", "mech", "deco", "safety"}

# 中文映射（总纲第 8 章）
DEFECT_CN = {"structural": "结构变形", "protection": "防护缺失", "installation": "安装不规范",
             "material": "材料质量", "safety": "安全隐患"}
STATUS_CN = {"pending": "待接收", "accepted": "已接收", "processing": "整改中", "review": "待审核", "closed": "已闭环"}

# Excel 填写支持中文标签或英文枚举码
LABEL_MAPS = {
    "defect_category": {"结构变形": "structural", "防护缺失": "protection", "安装不规范": "installation",
                        "材料质量": "material", "安全隐患": "safety"},
    "risk_level": {"高": "high", "高险": "high", "中": "medium", "中险": "medium", "低": "low", "低险": "low"},
    "types": {"房建": "building", "道路": "road", "隧道": "tunnel", "景观": "landscape"},
    "phases": {"基坑": "foundation", "主体结构": "structure", "机电安装": "mep", "装修": "decoration"},
}

TEMPLATE_HEADERS = ["名称*", "检查要点（分号分隔）*", "依据条款*", "问题类别*", "风险等级*",
                    "适用项目类型（逗号分隔）*", "适用建设阶段（逗号分隔）*"]


def _client_ip(request: Request):
    return request.client.host if request.client else None


# ---------- 知识库（org 隔离，写操作需 knowledge:write） ----------

class ItemIn(BaseModel):
    name: str
    check_points: list[str]
    basis: str
    defect_category: str
    risk_level: str
    applicable_types: list[str]
    applicable_phases: list[str]
    enabled: bool = True

    @field_validator("name", "basis")
    @classmethod
    def not_empty(cls, v: str):
        if not v or not v.strip():
            raise ValueError("不能为空")
        return v.strip()

    @field_validator("defect_category")
    @classmethod
    def valid_category(cls, v: str):
        if v not in DEFECT_CATEGORIES:
            raise ValueError("问题类别枚举不合法")
        return v

    @field_validator("risk_level")
    @classmethod
    def valid_risk(cls, v: str):
        if v not in RISK_LEVELS:
            raise ValueError("风险等级枚举不合法")
        return v

    @field_validator("applicable_types")
    @classmethod
    def valid_types(cls, v: list[str]):
        bad = [x for x in v if x not in PROJECT_TYPES]
        if bad:
            raise ValueError(f"项目类型枚举不合法: {bad}")
        return v

    @field_validator("applicable_phases")
    @classmethod
    def valid_phases(cls, v: list[str]):
        bad = [x for x in v if x not in PHASES]
        if bad:
            raise ValueError(f"建设阶段枚举不合法: {bad}")
        return v


class ItemUpdateIn(ItemIn):
    expected_version: int = Field(ge=1)


def _normalize_enum(value: str, kind: str) -> str | None:
    """Excel 单元格枚举值归一化：接受英文码或中文标签"""
    value = (value or "").strip()
    if not value:
        return None
    if kind in ("defect_category", "risk_level"):
        if value in (DEFECT_CATEGORIES if kind == "defect_category" else RISK_LEVELS):
            return value
        return LABEL_MAPS[kind].get(value)
    if kind == "types":
        return value if value in PROJECT_TYPES else LABEL_MAPS["types"].get(value)
    if kind == "phases":
        return value if value in PHASES else LABEL_MAPS["phases"].get(value)
    return None


def _split_multi(value: str, kind: str) -> list[str]:
    """逗号分隔多值列解析，逐个归一化"""
    result = []
    for part in (value or "").replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        normalized = _normalize_enum(part, kind)
        result.append(normalized if normalized else part)
    return result


@router.get("/inspection-items/template")
def download_template(scope: ProjectScope = Depends(require_perm("knowledge:write"))):
    """动态生成 xlsx 导入模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "检查项"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["示例检查项名称", "要点一；要点二；要点三", "JGJ 59-2011《建筑施工安全检查标准》表3.13",
               "safety", "high", "building,road", "foundation,structure"])
    for i, _ in enumerate(TEMPLATE_HEADERS, start=1):
        ws.column_dimensions[chr(64 + i)].width = 26
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''%E6%A3%80%E6%9F%A5%E9%A1%B9%E5%AF%BC%E5%85%A5%E6%A8%A1%E6%9D%BF.xlsx"},
    )


@router.post("/inspection-items/import")
async def import_items(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    scope: ProjectScope = Depends(require_perm("knowledge:write")),
):
    """xlsx 批量导入：逐行校验，合法入库（挂当前组织），非法记录行号与原因"""
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True)
    except Exception:
        raise BizError(1003, "文件解析失败，请使用标准模板")
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
    success_count = 0
    fail_details = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue  # 空行跳过
        def cell(pos: int) -> str:
            return str(row[pos]).strip() if pos < len(row) and row[pos] is not None else ""
        name, points_raw, basis, category_raw, risk_raw, types_raw, phases_raw = (
            cell(0), cell(1), cell(2), cell(3), cell(4), cell(5), cell(6))
        if not name:
            fail_details.append({"row": idx, "reason": "名称必填"})
            continue
        if not basis:
            fail_details.append({"row": idx, "reason": "依据条款必填"})
            continue
        category = _normalize_enum(category_raw, "defect_category")
        if not category:
            fail_details.append({"row": idx, "reason": f"问题类别不合法: {category_raw}"})
            continue
        risk = _normalize_enum(risk_raw, "risk_level")
        if not risk:
            fail_details.append({"row": idx, "reason": f"风险等级不合法: {risk_raw}"})
            continue
        types = _split_multi(types_raw, "types")
        bad_types = [t for t in types if t not in PROJECT_TYPES]
        if not types or bad_types:
            fail_details.append({"row": idx, "reason": f"适用项目类型不合法: {types_raw}"})
            continue
        phases = _split_multi(phases_raw, "phases")
        bad_phases = [p for p in phases if p not in PHASES]
        if not phases or bad_phases:
            fail_details.append({"row": idx, "reason": f"适用建设阶段不合法: {phases_raw}"})
            continue
        points = [p.strip() for p in points_raw.replace("；", ";").split(";") if p.strip()]
        if not points:
            fail_details.append({"row": idx, "reason": "检查要点必填"})
            continue
        db.add(InspectionItem(
            org_id=scope.org_id,
            name=name,
            check_points=json.dumps(points, ensure_ascii=False),
            basis=basis,
            defect_category=category,
            risk_level=risk,
            applicable_types=json.dumps(types),
            applicable_phases=json.dumps(phases),
            updated_by=scope.user.id,
            enabled=False,
            publication="draft",
        ))
        success_count += 1
    write_audit(db, "knowledge_import", actor_id=scope.user.id, target_type="inspection_item",
                result="success", org_id=scope.org_id, project_id=scope.project_id,
                after={"success_count": success_count}, commit=False)
    db.commit()
    return ok({"success_count": success_count, "fail_details": fail_details})


@router.get("/inspection-items")
def admin_list_items(
    keyword: str | None = None,
    project_type: str | None = None,
    phase: str | None = None,
    defect_category: str | None = None,
    risk_level: str | None = None,
    enabled: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    scope: ProjectScope = Depends(require_perm("knowledge:write")),
):
    """admin 列表（org 隔离）：enabled=all 返回全部；支持问题类别/风险等级筛选"""
    if defect_category and defect_category not in DEFECT_CATEGORIES:
        raise BizError(1000, "问题类别枚举不合法")
    if risk_level and risk_level not in RISK_LEVELS:
        raise BizError(1000, "风险等级枚举不合法")
    q = build_item_query(db, keyword, project_type, phase, enabled, scope.org_id)
    if defect_category:
        q = q.filter(InspectionItem.defect_category == defect_category)
    if risk_level:
        q = q.filter(InspectionItem.risk_level == risk_level)
    total = q.count()
    items = q.order_by(InspectionItem.id).offset((page - 1) * page_size).limit(page_size).all()
    return ok({"list": [item_to_dict(i) for i in items], "total": total, "page": page, "page_size": page_size})


@router.post("/inspection-items")
def create_item(body: ItemIn, db: Session = Depends(get_db),
                scope: ProjectScope = Depends(require_perm("knowledge:write"))):
    item = InspectionItem(
        org_id=scope.org_id,
        name=body.name,
        check_points=json.dumps(body.check_points, ensure_ascii=False),
        basis=body.basis,
        defect_category=body.defect_category,
        risk_level=body.risk_level,
        applicable_types=json.dumps(body.applicable_types),
        applicable_phases=json.dumps(body.applicable_phases),
        updated_by=scope.user.id,
        enabled=False,
        publication="draft",
    )
    db.add(item)
    db.flush()
    write_audit(db, "knowledge_create", actor_id=scope.user.id, target_type="inspection_item",
                target_id=item.id, org_id=scope.org_id, project_id=scope.project_id, commit=False)
    db.commit()
    db.refresh(item)
    return ok(item_to_dict(item))


@router.put("/inspection-items/{item_id}")
def update_item(item_id: int, body: ItemUpdateIn, db: Session = Depends(get_db),
                scope: ProjectScope = Depends(require_perm("knowledge:write"))):
    item = db.get(InspectionItem, item_id)
    # 平台内置模板（org_id IS NULL）不可被组织改动；他组织条目不可见
    if item is None or (item.org_id is None) or item.org_id != scope.org_id:
        raise BizError(1004, "检查项不存在")
    from ..services.operations import revise_item
    revise_item(db, item, body.expected_version, dict(
        name=body.name, check_points=json.dumps(body.check_points, ensure_ascii=False),
        basis=body.basis, defect_category=body.defect_category, risk_level=body.risk_level,
        applicable_types=json.dumps(body.applicable_types), applicable_phases=json.dumps(body.applicable_phases),
        enabled=False, publication="draft", approved_by=None, effective_from=None,
        effective_until=None, updated_by=scope.user.id))
    write_audit(db, "knowledge_update", actor_id=scope.user.id, target_type="inspection_item",
                target_id=item.id, org_id=scope.org_id, project_id=scope.project_id, commit=False)
    db.commit()
    db.refresh(item)
    return ok(item_to_dict(item))


@router.delete("/inspection-items/{item_id}")
def delete_item(item_id: int, expected_version: int = Query(ge=1), db: Session = Depends(get_db),
                scope: ProjectScope = Depends(require_perm("knowledge:write"))):
    item = db.get(InspectionItem, item_id)
    if item is None or (item.org_id is None) or item.org_id != scope.org_id:
        raise BizError(1004, "检查项不存在")
    # 注：检查项属组织级知识库，删除前需确认组织内任一项目均未引用，故刻意不加 project 过滤（不涉及跨项目数据可见）。
    if db.query(InspectionRecord).filter(InspectionRecord.item_id == item_id).count() > 0:
        raise BizError(400, "该检查项已被巡查记录引用，请停用而非删除")
    from ..services.operations import revise_item
    revise_item(db, item, expected_version, dict(enabled=False, publication="withdrawn", updated_by=scope.user.id))
    write_audit(db, "knowledge_delete", actor_id=scope.user.id, target_type="inspection_item",
                target_id=item.id, org_id=scope.org_id, project_id=scope.project_id, commit=False)
    db.commit()
    return ok(None)


# ---------- 成员管理（当前项目，需 member:write） ----------

class UserCreate(BaseModel):
    username: str
    name: str
    role: str
    specialty: str | None = None
    phone: str | None = None
    org_unit: str | None = None

    @field_validator("role")
    @classmethod
    def valid_role(cls, v: str):
        if v not in MEMBER_ROLES:
            raise ValueError("角色枚举不合法")
        return v

    @field_validator("specialty")
    @classmethod
    def valid_specialty(cls, v: str | None):
        if v and v not in SPECIALTIES:
            raise ValueError("专业枚举不合法")
        return v


class UserUpdate(BaseModel):
    name: str | None = None
    phone: str | None = None
    role: str | None = None
    specialty: str | None = None
    org_unit: str | None = None
    active: bool | None = None

    @field_validator("role")
    @classmethod
    def valid_role(cls, v: str | None):
        if v and v not in MEMBER_ROLES:
            raise ValueError("角色枚举不合法")
        return v

    @field_validator("specialty")
    @classmethod
    def valid_specialty(cls, v: str | None):
        if v and v not in SPECIALTIES:
            raise ValueError("专业枚举不合法")
        return v


def _member_to_dict(db: Session, u: User, members: list[ProjectMember], project_id: int) -> dict:
    open_count = (
        db.query(RectifyOrder)
        .filter(RectifyOrder.assignee_id == u.id, RectifyOrder.project_id == project_id,
                RectifyOrder.status != "closed")
        .count()
    )
    m0 = members[0] if members else None
    return {
        "id": u.id,
        "username": u.username,
        "name": u.name,
        "phone": u.phone,
        "roles": [m.role for m in members],
        "specialty": m0.specialty if m0 else None,
        "org_unit": m0.org_unit if m0 else None,
        "active": u.active and any(m.active for m in members) if members else u.active,
        "project_id": project_id,
        "open_count": open_count,
    }


@router.get("/users")
def list_users(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
               db: Session = Depends(get_db),
               scope: ProjectScope = Depends(require_perm("member:write"))):
    """当前项目成员分页列表（按用户聚合多角色，含未闭环工单数；P05 §2.4 数据库层分页）"""
    # 先按 user_id 去重在数据库层分页，再取页内用户的成员关系（页内结果有界：≤ page_size × 角色数）
    ids_q = (
        db.query(ProjectMember.user_id)
        .filter(ProjectMember.project_id == scope.project_id)
        .distinct()
    )
    total = ids_q.count()
    page_ids = [
        r[0] for r in (
            ids_q.order_by(ProjectMember.user_id)
            .offset((page - 1) * page_size).limit(page_size).all()
        )
    ]
    member_rows = (
        db.query(ProjectMember)
        .filter(ProjectMember.project_id == scope.project_id, ProjectMember.user_id.in_(page_ids))
        .all()  # 仅当前页用户的成员关系行，结果集有界
    )
    by_user: dict[int, list[ProjectMember]] = {}
    for m in member_rows:
        by_user.setdefault(m.user_id, []).append(m)
    result = []
    for uid in page_ids:
        u = db.get(User, uid)
        if u is None:
            continue
        result.append(_member_to_dict(db, u, by_user.get(uid, []), scope.project_id))
    return ok({"list": result, "total": total, "page": page, "page_size": page_size})


@router.post("/users")
def create_user(body: UserCreate, request: Request, db: Session = Depends(get_db),
                scope: ProjectScope = Depends(require_perm("member:write"))):
    """新增成员：创建用户（挂当前组织）+ 当前项目成员关系；系统生成一次性强密码"""
    if db.query(User).filter(User.username == body.username).first():
        raise BizError(1021, "用户名已存在")
    initial_password = generate_strong_password()
    new_user = User(
        username=body.username,
        name=body.name,
        phone=body.phone,
        org_id=scope.org_id,
        password_hash=hash_password(initial_password),
        active=True,
        must_change_password=True,
        password_updated_at=now_utc(),
    )
    db.add(new_user)
    db.flush()
    db.add(ProjectMember(
        project_id=scope.project_id,
        user_id=new_user.id,
        role=body.role,
        specialty=body.specialty,
        org_unit=body.org_unit,
        is_default=True,
        active=True,
    ))
    write_audit(db, "user_created", actor_id=scope.user.id, target_type="user", target_id=new_user.id,
                after={"username": new_user.username, "role": body.role}, org_id=scope.org_id,
                project_id=scope.project_id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(new_user)
    data = _member_to_dict(db, new_user, [db.query(ProjectMember).filter(
        ProjectMember.user_id == new_user.id, ProjectMember.project_id == scope.project_id).first()],
        scope.project_id)
    data["initial_password"] = initial_password
    return {"code": 0, "msg": "创建成功，初始密码仅显示一次，请立即复制交予本人", "data": data}


@router.post("/users/{user_id}/reset-password")
def reset_password(user_id: int, request: Request, db: Session = Depends(get_db),
                   scope: ProjectScope = Depends(require_perm("member:write"))):
    """重置成员密码：目标须为本组织用户；随机强密码 + 强制改密 + 吊销旧令牌"""
    target = db.get(User, user_id)
    if target is None or target.org_id != scope.org_id:
        raise BizError(1022, "用户不存在")
    initial_password = generate_strong_password()
    target.password_hash = hash_password(initial_password)
    target.must_change_password = True
    target.password_updated_at = now_utc()
    target.token_version += 1
    target.failed_login_count = 0
    target.locked_until = None
    revoke_all_refresh(db, target.id)
    write_audit(db, "password_reset", actor_id=scope.user.id, target_type="user", target_id=target.id,
                org_id=scope.org_id, project_id=scope.project_id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    return {"code": 0, "msg": "重置成功，初始密码仅显示一次，请立即复制交予本人",
            "data": {"id": target.id, "username": target.username, "initial_password": initial_password}}


@router.put("/users/{user_id}")
def update_user(user_id: int, body: UserUpdate, request: Request, db: Session = Depends(get_db),
                scope: ProjectScope = Depends(require_perm("member:write"))):
    """编辑本项目成员：改用户基本信息 + 本项目成员角色/专业/单位/启用状态"""
    target = db.get(User, user_id)
    if target is None or target.org_id != scope.org_id:
        raise BizError(1022, "用户不存在")
    members = (
        db.query(ProjectMember)
        .filter(ProjectMember.user_id == user_id, ProjectMember.project_id == scope.project_id)
        .all()  # 结果集有界：单用户单项目角色行（≤ 角色枚举基数）
    )
    if not members:
        raise BizError(1022, "该用户不是本项目成员")
    if body.name is not None:
        target.name = body.name
    if body.phone is not None:
        target.phone = body.phone
    primary = members[0]
    old_role = primary.role
    role_changed = body.role is not None and body.role != old_role
    disabled = body.active is False
    if body.role is not None:
        primary.role = body.role
    if body.specialty is not None:
        primary.specialty = body.specialty
    if body.org_unit is not None:
        primary.org_unit = body.org_unit
    if body.active is not None:
        for m in members:
            m.active = body.active
    ip = _client_ip(request)
    ua = request.headers.get("user-agent")
    # 成员停用或角色变更 → 令牌立即失效（token_version 是账号级）
    if role_changed or disabled:
        from ..services.operations import ensure_handoff
        for member in members:
            ensure_handoff(db,scope.project_id,target.id,member.role if disabled else old_role)
        target.token_version += 1
        revoke_all_refresh(db, target.id)
    if disabled:
        write_audit(db, "member_disabled", actor_id=scope.user.id, target_type="user", target_id=target.id,
                    org_id=scope.org_id, project_id=scope.project_id, ip=ip, user_agent=ua, commit=False)
    if role_changed:
        write_audit(db, "member_role_changed", actor_id=scope.user.id, target_type="user", target_id=target.id,
                    after={"role": primary.role}, org_id=scope.org_id, project_id=scope.project_id,
                    ip=ip, user_agent=ua, commit=False)
    db.commit()
    members = (
        db.query(ProjectMember)
        .filter(ProjectMember.user_id == user_id, ProjectMember.project_id == scope.project_id)
        .all()  # 结果集有界：单用户单项目角色行（≤ 角色枚举基数）
    )
    db.refresh(target)
    return ok(_member_to_dict(db, target, members, scope.project_id))


# ---------- 成员关系 CRUD（/api/admin/members） ----------

class MemberAdd(BaseModel):
    user_id: int
    role: str
    specialty: str | None = None
    org_unit: str | None = None
    is_default: bool = False

    @field_validator("role")
    @classmethod
    def valid_role(cls, v: str):
        if v not in MEMBER_ROLES:
            raise ValueError("角色枚举不合法")
        return v


class MemberUpdate(BaseModel):
    role: str | None = None
    specialty: str | None = None
    org_unit: str | None = None
    active: bool | None = None

    @field_validator("role")
    @classmethod
    def valid_role(cls, v: str | None):
        if v and v not in MEMBER_ROLES:
            raise ValueError("角色枚举不合法")
        return v


@router.get("/members")
def list_members(db: Session = Depends(get_db),
                 scope: ProjectScope = Depends(require_perm("member:write"))):
    """本项目成员关系明细（一角色一行）"""
    # 结果集有界：单项目成员角色行（成员管理页一次性展示，量级 < 数百）
    rows = db.query(ProjectMember).filter(ProjectMember.project_id == scope.project_id).order_by(ProjectMember.id).all()
    result = []
    for m in rows:
        u = db.get(User, m.user_id)
        result.append({"member_id": m.id, "user_id": m.user_id, "username": u.username if u else None,
                       "name": u.name if u else None, "role": m.role, "specialty": m.specialty,
                       "org_unit": m.org_unit, "is_default": m.is_default, "active": m.active})
    return ok(result)


@router.post("/members")
def add_member(body: MemberAdd, request: Request, db: Session = Depends(get_db),
               scope: ProjectScope = Depends(require_perm("member:write"))):
    """把本组织已有用户加入当前项目并授予角色"""
    target = db.get(User, body.user_id)
    if target is None or target.org_id != scope.org_id:
        raise BizError(1022, "用户不存在或不属于本组织")
    exists = (
        db.query(ProjectMember)
        .filter(ProjectMember.project_id == scope.project_id, ProjectMember.user_id == body.user_id,
                ProjectMember.role == body.role)
        .first()
    )
    if exists:
        raise BizError(1021, "该成员已具备此角色")
    member = ProjectMember(
        project_id=scope.project_id, user_id=body.user_id, role=body.role,
        specialty=body.specialty, org_unit=body.org_unit, is_default=body.is_default, active=True,
    )
    db.add(member)
    write_audit(db, "member_added", actor_id=scope.user.id, target_type="user", target_id=body.user_id,
                after={"role": body.role}, org_id=scope.org_id, project_id=scope.project_id,
                ip=_client_ip(request), user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(member)
    return ok({"member_id": member.id})


@router.put("/members/{member_id}")
def update_member(member_id: int, body: MemberUpdate, request: Request, db: Session = Depends(get_db),
                  scope: ProjectScope = Depends(require_perm("member:write"))):
    """修改成员关系：角色 / 专业 / 单位 / 启用；跨项目成员关系不可越权修改"""
    m = db.get(ProjectMember, member_id)
    if m is None or m.project_id != scope.project_id:
        raise BizError(1022, "成员关系不存在")
    if body.active is False or (body.role is not None and body.role != m.role):
        from ..services.operations import ensure_handoff
        ensure_handoff(db,scope.project_id,m.user_id,m.role)
    if body.role is not None:
        m.role = body.role
    if body.specialty is not None:
        m.specialty = body.specialty
    if body.org_unit is not None:
        m.org_unit = body.org_unit
    if body.active is not None:
        m.active = body.active
    write_audit(db, "member_updated", actor_id=scope.user.id, target_type="user", target_id=m.user_id,
                after={"role": m.role, "active": m.active}, org_id=scope.org_id, project_id=scope.project_id,
                ip=_client_ip(request), user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    return ok({"member_id": m.id})


@router.delete("/members/{member_id}")
def remove_member(member_id: int, request: Request, db: Session = Depends(get_db),
                  scope: ProjectScope = Depends(require_perm("member:write"))):
    """移除成员关系（跨项目不可越权删除）"""
    m = db.get(ProjectMember, member_id)
    if m is None or m.project_id != scope.project_id:
        raise BizError(1022, "成员关系不存在")
    write_audit(db, "member_removed", actor_id=scope.user.id, target_type="user", target_id=m.user_id,
                org_id=scope.org_id, project_id=scope.project_id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    from ..services.operations import ensure_handoff
    ensure_handoff(db,scope.project_id,m.user_id,m.role)
    db.delete(m)
    db.commit()
    return ok(None)


# ---------- 项目信息 ----------

class ProjectUpdate(BaseModel):
    name: str | None = None
    project_type: str | None = None
    phase: str | None = None
    address: str | None = None
    status: str | None = None

    @field_validator("project_type")
    @classmethod
    def valid_type(cls, v: str | None):
        if v and v not in PROJECT_TYPES:
            raise ValueError("项目类型枚举不合法")
        return v

    @field_validator("phase")
    @classmethod
    def valid_phase(cls, v: str | None):
        if v and v not in PHASES:
            raise ValueError("建设阶段枚举不合法")
        return v

    @field_validator("status")
    @classmethod
    def valid_status(cls, v: str | None):
        if v and v not in ("active", "paused", "closed"):
            raise ValueError("项目状态枚举不合法")
        return v


def _project_dict(p: Project) -> dict:
    return {"id": p.id, "org_id": p.org_id, "name": p.name, "code": p.code,
            "project_type": p.project_type, "phase": p.phase, "address": p.address,
            "status": p.status}


@router.get("/project")
def get_project(db: Session = Depends(get_db), scope: ProjectScope = Depends(require_perm("member:write"))):
    """当前项目详情（不再锁定 id 最小项目，改为操作 scope.project_id）"""
    project = db.get(Project, scope.project_id)
    if project is None:
        raise BizError(1023, "项目不存在")
    return ok(_project_dict(project))


@router.put("/project")
def update_project(body: ProjectUpdate, request: Request, db: Session = Depends(get_db),
                   scope: ProjectScope = Depends(require_perm("member:write"))):
    """更新当前项目（scope.project_id），不影响其他项目"""
    project = db.get(Project, scope.project_id)
    if project is None:
        raise BizError(1023, "项目不存在")
    if body.name is not None:
        project.name = body.name
    if body.project_type is not None:
        project.project_type = body.project_type
    if body.phase is not None:
        project.phase = body.phase
    if body.address is not None:
        project.address = body.address
    if body.status is not None:
        project.status = body.status
    write_audit(db, "project_updated", actor_id=scope.user.id, target_type="project", target_id=project.id,
                org_id=scope.org_id, project_id=project.id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(project)
    return ok(_project_dict(project))


# ---------- 组织级项目管理（org:manage） ----------

class ProjectCreate(BaseModel):
    name: str
    project_type: str
    phase: str
    code: str | None = None
    address: str | None = None

    @field_validator("project_type")
    @classmethod
    def valid_type(cls, v: str):
        if v not in PROJECT_TYPES:
            raise ValueError("项目类型枚举不合法")
        return v

    @field_validator("phase")
    @classmethod
    def valid_phase(cls, v: str):
        if v not in PHASES:
            raise ValueError("建设阶段枚举不合法")
        return v


@router.get("/projects")
def list_projects(db: Session = Depends(get_db), scope: ProjectScope = Depends(require_perm("org:manage"))):
    """组织管理员：本组织全部项目"""
    # 结果集有界：单组织项目数（项目下拉/管理页一次性展示）
    projects = db.query(Project).filter(Project.org_id == scope.org_id).order_by(Project.id).all()
    return ok([_project_dict(p) for p in projects])


@router.post("/projects")
def create_project(body: ProjectCreate, request: Request, db: Session = Depends(get_db),
                   scope: ProjectScope = Depends(require_perm("org:manage"))):
    """组织管理员：在本组织新建项目"""
    project = Project(
        org_id=scope.org_id, name=body.name, code=body.code,
        project_type=body.project_type, phase=body.phase, address=body.address, status="active",
    )
    db.add(project)
    db.flush()
    write_audit(db, "project_created", actor_id=scope.user.id, target_type="project", target_id=project.id,
                after={"name": project.name}, org_id=scope.org_id, project_id=project.id,
                ip=_client_ip(request), user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(project)
    return ok(_project_dict(project))


@router.put("/projects/{project_id}")
def update_org_project(project_id: int, body: ProjectUpdate, request: Request, db: Session = Depends(get_db),
                       scope: ProjectScope = Depends(require_perm("org:manage"))):
    """组织管理员：更新本组织内指定项目（跨组织不可越权）"""
    project = db.get(Project, project_id)
    if project is None or project.org_id != scope.org_id:
        raise BizError(1023, "项目不存在")
    if body.name is not None:
        project.name = body.name
    if body.project_type is not None:
        project.project_type = body.project_type
    if body.phase is not None:
        project.phase = body.phase
    if body.address is not None:
        project.address = body.address
    if body.status is not None:
        project.status = body.status
    write_audit(db, "project_updated", actor_id=scope.user.id, target_type="project", target_id=project.id,
                org_id=scope.org_id, project_id=project.id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(project)
    return ok(_project_dict(project))


# ---------- 工单监督（按当前项目聚合） ----------

class UrgeBatchIn(BaseModel):
    """批量催办入参（P06 §2.3.2）：最多 50 条"""
    order_ids: list[int]

    @field_validator("order_ids")
    @classmethod
    def valid_ids(cls, v: list[int]):
        if not v:
            raise ValueError("order_ids 不能为空")
        if len(v) > 50:
            raise ValueError("单次批量催办最多 50 条")
        if len(set(v)) != len(v):
            raise ValueError("order_ids 不得重复")
        return v


@router.post("/orders/urge-batch")
def urge_batch(body: UrgeBatchIn, db: Session = Depends(get_db),
               scope: ProjectScope = Depends(require_perm("order:urge"))):
    """批量催办（P06 §2.3.2）：逐单复用催办逻辑；闭环单/限流窗口内的单返回跳过原因，不阻断其余"""
    from .orders import _get_scoped_order
    from ..services.order_flow import log_urge

    urged: list[dict] = []
    skipped: list[dict] = []
    for oid in body.order_ids:
        try:
            order = _get_scoped_order(oid, db, scope)
            if order.status == "closed":
                skipped.append({"order_id": oid, "reason": "工单已闭环，无需催办"})
                continue
            if order.status == "cancelled":
                skipped.append({"order_id": oid, "reason": "工单已作废"})
                continue
            interval_hours = resolve_urge_interval_hours(db, order)
            last_urge = (
                db.query(OrderLog)
                .filter(OrderLog.order_id == order.id, OrderLog.action == "urge")
                .order_by(OrderLog.id.desc())
                .first()
            )
            if last_urge and last_urge.created_at and \
                    now_utc() - as_utc(last_urge.created_at) < timedelta(hours=interval_hours):
                skipped.append({"order_id": oid, "reason": f"{interval_hours}小时限流窗口内已催办"})
                continue
            log_urge(order, scope.user, db)  # 每单独立事务提交 + 提交后即时投递
            urged.append({"order_id": oid, "order_no": order.order_no})
        except BizError as e:
            skipped.append({"order_id": oid, "reason": e.msg})
    return ok({"urged": urged, "skipped": skipped})


@router.get("/orders")
def admin_list_orders(
    status: str | None = None,
    overdue: bool | None = None,
    warn: bool | None = None,
    assignee_id: int | None = None,
    sort: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    scope: ProjectScope = Depends(project_scope),
):
    """监督列表（P05 §2.4：数据库层分页 + overdue SQL 化；P06：即将超期筛选 + 超期时长排序）"""
    if not scope.can_read_all:
        raise forbidden("无权限")
    if sort is not None and sort not in ("overdue_asc", "overdue_desc"):
        raise BizError(1000, "sort 参数不合法（overdue_asc/overdue_desc）")
    q = db.query(RectifyOrder).filter(RectifyOrder.project_id == scope.project_id)
    if status:
        q = q.filter(RectifyOrder.status == status)
    if assignee_id:
        q = q.filter(RectifyOrder.assignee_id == assignee_id)
    now = now_utc()
    if overdue:
        # 超期过滤 SQL 化：未闭环未作废且截止时间已过（与 order_to_dict._is_overdue 口径一致）
        q = q.filter(
            RectifyOrder.status.notin_(("closed", "cancelled")),
            RectifyOrder.deadline < now,
        )
    if warn:
        # 即将超期（P06 §2.5.1）：已进入临期提醒窗口（now >= warn_at）且尚未超期、未闭环未作废
        q = q.filter(
            RectifyOrder.status.notin_(("closed", "cancelled")),
            RectifyOrder.warn_at.isnot(None),
            RectifyOrder.warn_at <= now,
            RectifyOrder.deadline > now,
        )
    total = q.count()
    if sort in ("overdue_asc", "overdue_desc"):
        # 超期时长排序：未闭环未作废的超期单按 deadline 升/降序（时长与 deadline 反序），其余单排后
        from sqlalchemy import case

        is_overdue_expr = (
            RectifyOrder.status.notin_(("closed", "cancelled")) & (RectifyOrder.deadline < now)
        )
        order_expr = case((is_overdue_expr, RectifyOrder.deadline), else_=None)
        direction = order_expr.asc() if sort == "overdue_asc" else order_expr.desc()
        q = q.order_by(direction.nulls_last(), RectifyOrder.id.desc())
    else:
        q = q.order_by(RectifyOrder.id.desc())
    orders = q.offset((page - 1) * page_size).limit(page_size).all()
    result = []
    for o in orders:
        d = order_to_dict(o, db)
        d["problem_location"] = o.problem_location
        d["overdue_hours"] = (
            round((now - as_utc(o.deadline)).total_seconds() / 3600, 1)
            if o.status not in ("closed", "cancelled") and as_utc(o.deadline) < now else None
        )
        result.append(d)
    return ok({"list": result, "total": total, "page": page, "page_size": page_size})


# ---------- 统计看板（按当前项目聚合） ----------

@router.get("/dashboard")
def dashboard(range: str = "7d", db: Session = Depends(get_db),
              scope: ProjectScope = Depends(project_scope)):
    """统计看板：强制 project 过滤；range=7d|30d|all（overdue_list 不受 range 限制）"""
    if not scope.can_read_all:
        raise forbidden("无权限")
    if range not in ("7d", "30d", "all"):
        raise BizError(1000, "range 参数不合法")
    now = now_utc()
    start = now - timedelta(days=7 if range == "7d" else 30 if range == "30d" else 36500)
    pid = scope.project_id

    # 巡查记录口径
    rec_q = db.query(InspectionRecord).filter(
        InspectionRecord.project_id == pid, InspectionRecord.created_at >= start)
    inspect_count = rec_q.count()
    abnormal_count = rec_q.filter(InspectionRecord.human_verdict == "abnormal").count()
    # 分组聚合结果有界（问题类别枚举基数 ≤ 5），非内存分页
    pie_rows = (
        rec_q.filter(InspectionRecord.human_verdict == "abnormal")
        .with_entities(InspectionRecord.defect_category, sa_func.count(InspectionRecord.id))
        .group_by(InspectionRecord.defect_category)
        .all()
    )
    problem_pie = [{"name": DEFECT_CN.get(cat, cat), "value": cnt} for cat, cnt in pie_rows]

    # 工单口径
    order_q = db.query(RectifyOrder).filter(
        RectifyOrder.project_id == pid, RectifyOrder.created_at >= start)
    # 分组聚合结果有界（工单状态枚举基数 ≤ 6），非内存分页
    funnel_rows = (
        order_q.with_entities(RectifyOrder.status, sa_func.count(RectifyOrder.id))
        .group_by(RectifyOrder.status)
        .all()
    )
    funnel_map = dict(funnel_rows)
    status_funnel = [{"name": STATUS_CN[s], "value": funnel_map.get(s, 0)}
                     for s in ("pending", "accepted", "processing", "review", "closed")]

    # AI 一致率（按项目）
    fb_q = db.query(AiFeedback).join(InspectionRecord, InspectionRecord.id == AiFeedback.record_id).filter(AiFeedback.project_id == pid, AiFeedback.created_at >= start, InspectionRecord.ai_mode == "real")
    fb_total = fb_q.count()
    fb_consistent = fb_q.filter(AiFeedback.consistent.is_(True)).count()
    ai_consistent_rate = round(fb_consistent / fb_total * 100, 1) if fb_total else 0

    # 平均整改时长（P05 §2.4：SQL 聚合，避免全量闭环工单载入内存）
    closed_q = order_q.filter(RectifyOrder.status == "closed", RectifyOrder.closed_at.isnot(None))
    dialect = db.bind.dialect.name if db.bind is not None else "sqlite"
    if dialect == "postgresql":
        dur = sa_func.extract("epoch", RectifyOrder.closed_at - RectifyOrder.created_at)
    elif dialect in ("mysql", "mariadb"):
        dur = sa_func.timestampdiff(text("SECOND"), RectifyOrder.created_at, RectifyOrder.closed_at)
    else:  # SQLite：julianday 差值（天）转秒
        dur = (sa_func.julianday(RectifyOrder.closed_at) - sa_func.julianday(RectifyOrder.created_at)) * 86400.0
    closed_cnt, closed_avg = closed_q.with_entities(
        sa_func.count(RectifyOrder.id), sa_func.avg(dur)).first()
    avg_rectify_hours = round(float(closed_avg) / 3600, 1) if closed_cnt and closed_avg is not None else 0

    # 超期治理（P05 §2.4）：看板只返回最紧急前 20 条 + overdue_total，
    # 完整列表走 /api/admin/orders?overdue=true 数据库分页
    overdue_q = db.query(RectifyOrder).filter(
        RectifyOrder.project_id == pid,
        RectifyOrder.status.notin_(("closed", "cancelled")),
        RectifyOrder.deadline < now,
    )
    overdue_total = overdue_q.count()
    overdue_orders = overdue_q.order_by(RectifyOrder.deadline.asc()).limit(20).all()
    overdue_list = []
    for o in overdue_orders:
        d = order_to_dict(o, db)
        d["overdue_hours"] = round((now - as_utc(o.deadline)).total_seconds() / 3600, 1)
        overdue_list.append(d)

    return ok({
        "problem_pie": problem_pie,
        "status_funnel": status_funnel,
        "metrics": {
            "inspect_count": inspect_count,
            "abnormal_count": abnormal_count,
            "ai_consistent_rate": ai_consistent_rate,
            "avg_rectify_hours": avg_rectify_hours,
            # 按期完成率（P06 §2.5.3，口径见 services/metrics.py；统计全量闭环单，不受 range 限制）
            "sla_on_time_rate": sla_on_time_rate(db, pid),
            # P1-3：口径说明——超期未闭环单不计入分母，避免"100% vs 超期清单"观感矛盾
            "sla_on_time_hint": (
                "按期完成率=已闭环工单中按期完成的比例；当前另有 "
                f"{overdue_total} 单未闭环且已超期，不计入分母"
                if overdue_total > 0
                else "按期完成率=已闭环工单中按期完成的比例"
            ),
        },
        "overdue_list": overdue_list,
        "overdue_total": overdue_total,
    })


# ---------- 通知发送记录（当前项目，可读全部角色；总纲 §5.7 / P06 §2.5.2） ----------

NOTIFY_STATUS = {"pending", "sent", "failed", "skipped"}
NOTIFY_CHANNEL = {"inapp", "wechat", "sms"}


def _notification_dict(n: Notification, db: Session) -> dict:
    user = db.get(User, n.user_id)
    return {
        "id": n.id,
        "user_id": n.user_id,
        "user_name": user.name if user else None,
        "title": n.title,
        "content": n.content,
        "biz_type": n.biz_type,
        "biz_id": n.biz_id,
        "channel": n.channel,
        "template_code": n.template_code,
        "status": n.status,
        "sent_at": fmt(n.sent_at),
        "retry_count": n.retry_count,
        "error": n.error,
        "read": n.read,
        "created_at": fmt(n.created_at),
    }


@router.get("/notifications")
def admin_list_notifications(
    status: str | None = None,
    channel: str | None = None,
    template_code: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    scope: ProjectScope = Depends(project_scope),
):
    """通知发送记录（数据库层分页）：按状态 / 通道 / 模板筛选；可读全部角色"""
    if not scope.can_read_all:
        raise forbidden("无权限")
    if status and status not in NOTIFY_STATUS:
        raise BizError(1000, "status 枚举不合法")
    if channel and channel not in NOTIFY_CHANNEL:
        raise BizError(1000, "channel 枚举不合法")
    q = db.query(Notification).filter(
        Notification.project_id == scope.project_id,
        Notification.org_id == scope.org_id,
    )
    if status:
        q = q.filter(Notification.status == status)
    if channel:
        q = q.filter(Notification.channel == channel)
    if template_code:
        q = q.filter(Notification.template_code == template_code)
    total = q.count()
    rows = (
        q.order_by(Notification.id.desc())
        .offset((page - 1) * page_size).limit(page_size).all()
    )
    return ok({"list": [_notification_dict(n, db) for n in rows],
               "total": total, "page": page, "page_size": page_size})


@router.post("/notifications/{notification_id}/resend")
def resend_notification(notification_id: int, request: Request, db: Session = Depends(get_db),
                        scope: ProjectScope = Depends(project_scope)):
    """手动重发失败/跳过的通知：立即走一次通道投递（失败仍记 error，不抛错）。
    仅 failed/skipped 可重发（sent/inapp 已送达无意义，pending 由调度器自然处理）。"""
    if not scope.can_read_all:
        raise forbidden("无权限")
    n = db.get(Notification, notification_id)
    if n is None or n.project_id != scope.project_id or n.org_id != scope.org_id:
        raise BizError(1024, "通知记录不存在")
    if n.status not in ("failed", "skipped"):
        raise BizError(400, "仅失败/跳过的通知可手动重发")
    from sqlalchemy import update
    from ..services.notify import dispatch_pending
    if db.execute(update(Notification).where(Notification.id==n.id,Notification.status.in_(("failed","skipped")))
        .values(status="pending",error=None,retry_count=0,lease_until=None)).rowcount != 1:
        raise BizError(1014,"通知状态已变化")
    db.commit()
    dispatch_pending(db)
    write_audit(db, "notification_resent", actor_id=scope.user.id, target_type="notification",
                target_id=notification_id, after={"status": n.status}, org_id=scope.org_id,
                project_id=scope.project_id, ip=_client_ip(request),
                user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    return ok(_notification_dict(n, db))


# ---------- SLA 时限规则（org 隔离，需 sla:write；总纲 §4.4 / P05 §2.3） ----------

ESCALATE_ROLES = {"project_admin", "org_admin"}


def _sla_rule_dict(r: SlaRule) -> dict:
    return {
        "id": r.id,
        "org_id": r.org_id,
        "project_id": r.project_id,  # None = 组织级默认
        "risk_level": r.risk_level,
        "defect_category": r.defect_category,  # None = 该风险等级全部类别
        "sla_hours": r.sla_hours,
        "warn_ratio": float(r.warn_ratio),
        "escalate_after_hours": r.escalate_after_hours,
        "escalate_to_role": r.escalate_to_role,
        "urge_interval_hours": r.urge_interval_hours,  # 催办限流窗口（P06 §2.3.1，空=默认 1 小时）
        "enabled": r.enabled,
        "created_at": fmt(r.created_at),
        "updated_at": fmt(r.updated_at),
    }


class SlaRuleIn(BaseModel):
    """SLA 规则入参：project_id 空=组织级默认；时限一旦写入工单即固化，改规则只影响新单"""
    project_id: int | None = None
    risk_level: str
    defect_category: str | None = None
    sla_hours: int
    warn_ratio: float = 0.75
    escalate_after_hours: int
    escalate_to_role: str | None = None
    urge_interval_hours: int | None = None   # 催办限流窗口（小时，空=默认 1 小时）
    enabled: bool = True

    @field_validator("risk_level")
    @classmethod
    def valid_risk(cls, v: str):
        if v not in RISK_LEVELS:
            raise ValueError("风险等级枚举不合法")
        return v

    @field_validator("defect_category")
    @classmethod
    def valid_category(cls, v: str | None):
        if v is not None and v not in DEFECT_CATEGORIES:
            raise ValueError("问题类别枚举不合法")
        return v

    @field_validator("urge_interval_hours")
    @classmethod
    def valid_urge_interval(cls, v: int | None):
        if v is not None and not 1 <= v <= 168:
            raise ValueError("催办限流窗口须为 1-168 小时")
        return v

    @field_validator("sla_hours")
    @classmethod
    def valid_hours(cls, v: int):
        if not 1 <= v <= 8760:
            raise ValueError("整改时限须为 1-8760 小时")
        return v

    @field_validator("warn_ratio")
    @classmethod
    def valid_ratio(cls, v: float):
        if not 0 < v <= 1:
            raise ValueError("临期提醒比例须为 (0, 1]")
        return v

    @field_validator("escalate_after_hours")
    @classmethod
    def valid_escalate(cls, v: int):
        if not 1 <= v <= 8760:
            raise ValueError("升级时限须为 1-8760 小时")
        return v

    @field_validator("escalate_to_role")
    @classmethod
    def valid_escalate_role(cls, v: str | None):
        if v is not None and v not in ESCALATE_ROLES:
            raise ValueError("升级角色只能为 project_admin/org_admin")
        return v


def _validate_rule_scope(db: Session, scope: ProjectScope, project_id: int | None) -> int | None:
    """规则挂载项目校验：须属本组织；非组织管理员只能配当前项目"""
    if project_id is None:
        return None
    if "org_admin" not in scope.effective_roles and project_id != scope.project_id:
        raise forbidden("仅组织管理员可为其他项目配置规则")
    project = db.get(Project, project_id)
    if project is None or project.org_id != scope.org_id:
        raise BizError(1023, "项目不存在")
    return project_id


def _dup_rule_exists(db: Session, scope: ProjectScope, body: SlaRuleIn, exclude_id: int | None = None) -> bool:
    """同 (org, project, risk_level, defect_category) 维度查重（表无唯一约束，业务层强制）"""
    q = db.query(SlaRule).filter(
        SlaRule.org_id == scope.org_id,
        SlaRule.project_id == body.project_id,
        SlaRule.risk_level == body.risk_level,
        SlaRule.defect_category == body.defect_category,
    )
    if exclude_id is not None:
        q = q.filter(SlaRule.id != exclude_id)
    return q.first() is not None


@router.get("/sla-rules")
def list_sla_rules(db: Session = Depends(get_db),
                   scope: ProjectScope = Depends(require_perm("sla:write"))):
    """SLA 规则列表：本组织全部（含项目级与组织级）；结果集有界（风险等级×类别×项目数组合）"""
    rules = (
        db.query(SlaRule)
        .filter(SlaRule.org_id == scope.org_id)
        .order_by(SlaRule.project_id.is_(None).desc(), SlaRule.project_id, SlaRule.risk_level, SlaRule.defect_category)
        .all()  # 结果集有界：规则为管理员手工配置，量级 < 数十
    )
    return ok([_sla_rule_dict(r) for r in rules])


@router.post("/sla-rules")
def create_sla_rule(body: SlaRuleIn, request: Request, db: Session = Depends(get_db),
                    scope: ProjectScope = Depends(require_perm("sla:write"))):
    """新增 SLA 规则（同维度唯一；只影响此后新建的工单，存量工单时限已固化）"""
    _validate_rule_scope(db, scope, body.project_id)
    if _dup_rule_exists(db, scope, body):
        raise BizError(1030, "该维度的 SLA 规则已存在，请直接修改原有规则")
    rule = SlaRule(
        org_id=scope.org_id,
        project_id=body.project_id,
        risk_level=body.risk_level,
        defect_category=body.defect_category,
        sla_hours=body.sla_hours,
        warn_ratio=body.warn_ratio,
        escalate_after_hours=body.escalate_after_hours,
        escalate_to_role=body.escalate_to_role,
        urge_interval_hours=body.urge_interval_hours,
        enabled=body.enabled,
    )
    db.add(rule)
    db.flush()
    write_audit(db, "sla_rule_created", actor_id=scope.user.id, target_type="sla_rule", target_id=rule.id,
                after=_sla_rule_dict(rule), org_id=scope.org_id, project_id=body.project_id,
                ip=_client_ip(request), user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(rule)
    return ok(_sla_rule_dict(rule))


@router.put("/sla-rules/{rule_id}")
def update_sla_rule(rule_id: int, body: SlaRuleIn, request: Request, db: Session = Depends(get_db),
                    scope: ProjectScope = Depends(require_perm("sla:write"))):
    """修改 SLA 规则（本组织规则；只影响此后新建的工单，存量工单时限已固化）"""
    rule = db.get(SlaRule, rule_id)
    if rule is None or rule.org_id != scope.org_id:
        raise BizError(1030, "SLA 规则不存在")
    _validate_rule_scope(db, scope, body.project_id)
    if _dup_rule_exists(db, scope, body, exclude_id=rule.id):
        raise BizError(1030, "该维度的 SLA 规则已存在，请直接修改原有规则")
    before = _sla_rule_dict(rule)
    rule.project_id = body.project_id
    rule.risk_level = body.risk_level
    rule.defect_category = body.defect_category
    rule.sla_hours = body.sla_hours
    rule.warn_ratio = body.warn_ratio
    rule.escalate_after_hours = body.escalate_after_hours
    rule.escalate_to_role = body.escalate_to_role
    rule.urge_interval_hours = body.urge_interval_hours
    rule.enabled = body.enabled
    write_audit(db, "sla_rule_updated", actor_id=scope.user.id, target_type="sla_rule", target_id=rule.id,
                before=before, after=_sla_rule_dict(rule), org_id=scope.org_id, project_id=body.project_id,
                ip=_client_ip(request), user_agent=request.headers.get("user-agent"), commit=False)
    db.commit()
    db.refresh(rule)
    return ok(_sla_rule_dict(rule))
